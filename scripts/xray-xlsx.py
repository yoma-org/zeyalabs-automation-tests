#!/usr/bin/env python3
"""Convert Xray Test Case Importer CSV files to formatted Excel (.xlsx).

Usage: python3 scripts/xray-xlsx.py <csv-dir> <out-dir>
Every <csv-dir>/**/*.csv becomes <out-dir>/**/*.xlsx with the same columns. Sheet 1 "Tests" is
the import sheet (Save As → CSV UTF-8 exports exactly it); the all-in-one file also gets "Modules".
Needs openpyxl (pip install openpyxl).
"""
import csv, sys, collections
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

WIDTH = {'Test ID': 8, 'Summary': 50, 'Test Type': 11, 'Priority': 10, 'Labels': 12,
         'Repository Path': 42, 'Description': 70, 'Preconditions': 45, 'Action': 50, 'Data': 30, 'Expected Result': 50}
HEAD_FILL = PatternFill('solid', fgColor='1F3864')
BAND = PatternFill('solid', fgColor='EEF3FA')
TOP = Border(top=Side(style='thin', color='8EA9DB'))


def convert(src: Path, dst: Path):
    rows = list(csv.reader(src.open(encoding='utf-8-sig', newline='')))
    header, body = rows[0], rows[1:]
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tests'
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical='center', wrap_text=True)
    prev, band = None, False
    for r in body:
        ws.append(r)
        n = ws.max_row
        new_test = r[0] != prev
        if new_test:
            band, prev = not band, r[0]
        for c in ws[n]:
            c.data_type = 's'  # never a formula, even if text starts with "="
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if band:
                c.fill = BAND
            if new_test:
                c.border = TOP
        ws.cell(n, 1).value = int(r[0]) if r[0].isdigit() else r[0]
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(h, 15)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions

    if dst.stem.endswith('-all'):
        m = wb.create_sheet('Modules')
        m.append(['Module (Repository Path)', 'Tests', 'Steps'])
        pi = header.index('Repository Path')
        tests, steps, mod_of = collections.OrderedDict(), collections.Counter(), {}
        for r in body:
            if r[pi]:
                mod_of[r[0]] = '/'.join(r[pi].split('/')[:2])
                tests[mod_of[r[0]]] = tests.get(mod_of[r[0]], 0) + 1
            steps[mod_of.get(r[0])] += 1
        for k, v in tests.items():
            m.append([k, v, steps[k]])
        m.append(['Total', sum(tests.values()), len(body)])
        for c in m[1] + m[m.max_row]:
            c.font = Font(bold=True)
        m.column_dimensions['A'].width = 45
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return len({r[0] for r in body}), len(body)


if __name__ == '__main__':
    src_dir, out_dir = Path(sys.argv[1]), Path(sys.argv[2])
    for f in sorted(src_dir.rglob('*.csv')):
        out = out_dir / f.relative_to(src_dir).with_suffix('.xlsx')
        t, s = convert(f, out)
        print(f'{out}: {t} tests, {s} rows')
